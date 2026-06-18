# app/products/router/tkp_generation.py

from io import BytesIO
from fastapi import APIRouter, HTTPException
from docxtpl import DocxTemplate
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook

router = APIRouter(prefix="/tkp_generation", tags=["TKP"])


# === TKP Generation Endpoints ===

@router.post("/")
async def tkp_generation(
        template_path: str,
        user_dict: dict
):
    if template_path.endswith(".docx"):
        doc = DocxTemplate(template_path)

        doc.render(user_dict)

        result_stream = BytesIO()
        doc.save(result_stream)
        result_stream.seek(0)

        return StreamingResponse(
            result_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": 'attachment; filename="generated_tkp.docx"'
            }
        )

    elif template_path.endswith(".xlsx"):
        workbook = load_workbook(template_path)

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        for key, value in user_dict.items():
                            placeholder = "{{ " + key + " }}"
                            cell.value = cell.value.replace(placeholder, str(value))

        result_stream = BytesIO()
        workbook.save(result_stream)
        result_stream.seek(0)

        return StreamingResponse(
            result_stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="generated_tkp.xlsx"'
            }
        )
    else:
        raise HTTPException(
            status_code=400,
            detail="Неподдерживаемый формат для файла"
        )
